"""
The downloadable .xlsx template for the user import, one per role.

National code (and class) columns are formatted as Text so Excel keeps
leading zeros; branch, grade and gender cells get dropdowns fed from a
hidden "lists" sheet (a range reference has no length limit, unlike an
inline list).
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.services.access import get_accessible_branches
from school.models import Branch, Grade
from staff.models import Staff

from .user_import import (
    BRANCH, GENDER, GRADE, MAX_ROWS, NATIONAL_CODE, ROLE_COLUMNS, SECTION,
)

TEXT_COLUMNS = {NATIONAL_CODE, SECTION}
WIDTHS = {NATIONAL_CODE: 16, BRANCH: 14, GRADE: 8, SECTION: 10, GENDER: 10}


def _branch_codes(acting_user):
    branches = (
        Branch.objects.filter(is_active=True)
        if acting_user.is_superuser
        else get_accessible_branches(acting_user).filter(is_active=True)
    )
    return list(branches.values_list("code", flat=True))


def _dropdown(sheet, lists, list_column, values, target_column, title):
    if not values:
        return
    letter = get_column_letter(list_column)
    for row, value in enumerate(values, start=1):
        lists.cell(row=row, column=list_column, value=value)

    validation = DataValidation(
        type="list",
        formula1=f"=lists!${letter}$1:${letter}${len(values)}",
        allow_blank=True,
        showErrorMessage=True,
        errorTitle="مقدار نامعتبر",
        error=f"«{title}» را از فهرست انتخاب کنید.",
    )
    target = get_column_letter(target_column)
    validation.add(f"{target}2:{target}{MAX_ROWS + 1}")
    sheet.add_data_validation(validation)


def build_template(role, acting_user):
    columns = ROLE_COLUMNS[role]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "کاربران"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="FAEAD2")
    for index, title in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        sheet.column_dimensions[get_column_letter(index)].width = WIDTHS.get(title, 18)

        if title in TEXT_COLUMNS:
            for row in range(2, MAX_ROWS + 2):
                sheet.cell(row=row, column=index).number_format = "@"

    lists = workbook.create_sheet("lists")
    lists.sheet_state = "hidden"

    if BRANCH in columns:
        _dropdown(sheet, lists, 1, _branch_codes(acting_user), columns.index(BRANCH) + 1, BRANCH)
    if GRADE in columns:
        levels = list(Grade.objects.filter(is_active=True).order_by("level").values_list("level", flat=True))
        _dropdown(sheet, lists, 2, levels, columns.index(GRADE) + 1, GRADE)
    if GENDER in columns:
        genders = [label for _, label in Staff.Gender.choices]
        _dropdown(sheet, lists, 3, genders, columns.index(GENDER) + 1, GENDER)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
