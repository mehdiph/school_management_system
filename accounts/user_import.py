"""
Bulk account creation from an .xlsx file (admin: «ورود گروهی از اکسل»).

    read_rows()   workbook -> ImportRow per non-empty data row (normalised)
    validate()    each row -> ok / error / skip, resolving branch, grade, class
    run_import()  both, then -- only if no row has an error -- creates every
                  "ok" row inside one transaction

Errors block the whole file (nothing is created); skips (the person already
has an account / an enrollment this year) are reported and left untouched.
Existing users are never modified, and classes, grades, branches and
academic years are never created here.
"""

import re
from collections import defaultdict
from dataclasses import dataclass, field

import jdatetime
from django.contrib.admin.models import ADDITION, LogEntry
from django.contrib.contenttypes.models import ContentType
from django.db import IntegrityError, transaction
from django.db.models import Q
from django.utils import timezone

from accounts.models import User
from accounts.validators import is_valid_national_code, normalize_national_code
from core.services.access import get_accessible_branches
from school.models import AcademicYear, Branch, Grade, SchoolClass
from staff.models import BranchAccess, Staff, TeacherAssignment, TeacherProfile
from student.models.student_enrollment import StudentEnrollment
from student.models.student_profile import StudentProfile
from supervisor.models.supervisor_profile import SupervisorProfile

MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_ROWS = 2000

STUDENT, TEACHER, SUPERVISOR = User.Roles.STUDENT, User.Roles.TEACHER, User.Roles.SUPERVISOR

ROLE_LABELS = {
    STUDENT: "دانش‌آموز",
    TEACHER: "معلم",
    SUPERVISOR: "پشتیبان",
}

FIRST_NAME, LAST_NAME, NATIONAL_CODE, BRANCH, GRADE, SECTION, GENDER = (
    "نام", "نام خانوادگی", "کد ملی", "کد شعبه", "پایه", "کلاس", "جنسیت",
)

ROLE_COLUMNS = {
    STUDENT: [FIRST_NAME, LAST_NAME, NATIONAL_CODE, BRANCH, GRADE, SECTION],
    TEACHER: [FIRST_NAME, LAST_NAME, NATIONAL_CODE, BRANCH, GENDER],
    SUPERVISOR: [FIRST_NAME, LAST_NAME, NATIONAL_CODE, BRANCH, GRADE],
}

#: Columns whose *value* may be blank (the header must still be there).
OPTIONAL_COLUMNS = {GENDER}

GENDER_VALUES = {
    "پسر": Staff.Gender.MALE,
    "مرد": Staff.Gender.MALE,
    "male": Staff.Gender.MALE,
    "دختر": Staff.Gender.FEMALE,
    "زن": Staff.Gender.FEMALE,
    "female": Staff.Gender.FEMALE,
}

#: "اول" .. "دوازدهم" -> 1 .. 12, the only non-numeric grade input accepted.
GRADE_ORDINALS = {
    name: level for level, name in enumerate(
        ("اول", "دوم", "سوم", "چهارم", "پنجم", "ششم",
         "هفتم", "هشتم", "نهم", "دهم", "یازدهم", "دوازدهم"),
        start=1,
    )
}

_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
_ARABIC_LETTERS = str.maketrans({"ي": "ی", "ك": "ک"})
_PERSIAN_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")


class FileRejected(Exception):
    """The file as a whole cannot be imported (wrong headers, too big, no current year...)."""


def fa_digits(value):
    return str(value).translate(_PERSIAN_DIGITS)


def normalize_text(value):
    """
    One spreadsheet cell -> clean text: ASCII digits, Persian ی/ک,
    trimmed, inner whitespace collapsed. ZWNJ (نیم‌فاصله) is kept: it is
    not whitespace (\\s does not match it).
    """

    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).translate(_DIGITS).translate(_ARABIC_LETTERS)
    return re.sub(r"\s+", " ", text).strip()


def parse_grade_level(value):
    """ "3", "۳", "3.0", "سوم", "پایه سوم" -> 3; anything else -> None."""

    text = normalize_text(value)
    text = re.sub(r"^پایه\s*", "", text)
    if text.isdigit():
        return int(text)
    return GRADE_ORDINALS.get(text)


def parse_gender(value):
    """'' for blank, a Staff.Gender value, or None when not recognised."""

    text = normalize_text(value).lower()
    if not text:
        return ""
    return GENDER_VALUES.get(text)


@dataclass
class ImportRow:
    number: int                     # Excel row number (header is row 1)
    values: dict                    # header -> normalised text
    national_code: str = ""
    status: str = "ok"              # ok / error / skip
    reasons: list = field(default_factory=list)
    # Resolved during validation:
    branch: Branch | None = None
    grade: Grade | None = None
    school_class: SchoolClass | None = None
    gender: str = ""

    @property
    def full_name(self):
        return f"{self.values.get(FIRST_NAME, '')} {self.values.get(LAST_NAME, '')}".strip()

    def error(self, reason):
        self.status = "error"
        self.reasons.append(reason)

    def skip(self, reason):
        self.status = "skip"
        self.reasons.append(reason)


@dataclass
class ImportResult:
    role: str
    rows: list
    created: list = field(default_factory=list)   # the new User objects
    committed: bool = False

    @property
    def errors(self):
        return [row for row in self.rows if row.status == "error"]

    @property
    def skipped(self):
        return [row for row in self.rows if row.status == "skip"]

    @property
    def problems(self):
        return [row for row in self.rows if row.status != "ok"]


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------


def _open_sheet(uploaded_file):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except Exception as exc:  # zip/xml/format errors all mean the same to the user
        raise FileRejected("فایل اکسل قابل خواندن نیست. یک فایل ‎.xlsx‎ معتبر بارگذاری کنید.") from exc
    return workbook, workbook.worksheets[0]


def _check_headers(header_row, role):
    headers = [normalize_text(cell) for cell in header_row]
    while headers and not headers[-1]:
        headers.pop()

    expected = ROLE_COLUMNS[role]
    # A trailing optional column may be left out entirely.
    required_prefix = expected
    while required_prefix and required_prefix[-1] in OPTIONAL_COLUMNS:
        required_prefix = required_prefix[:-1]

    if headers not in (expected, required_prefix):
        raise FileRejected(
            f"سرستون‌های فایل با نقش «{ROLE_LABELS[role]}» مطابقت ندارد. "
            f"سرستون‌های لازم به این ترتیب: {'، '.join(expected)}. "
            f"سرستون‌های فایل: {'، '.join(h or '(خالی)' for h in headers) or '(هیچ)'}."
        )


def read_rows(uploaded_file, role):
    workbook, sheet = _open_sheet(uploaded_file)
    try:
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise FileRejected("فایل خالی است.")
        _check_headers(header, role)

        columns = ROLE_COLUMNS[role]
        result = []
        for number, cells in enumerate(rows, start=2):
            cells = list(cells[:len(columns)]) + [None] * (len(columns) - len(cells))
            values = {
                column: normalize_text(cell) for column, cell in zip(columns, cells)
            }
            if not any(values.values()):
                continue  # fully empty row
            if len(result) == MAX_ROWS:
                raise FileRejected(f"فایل بیش از {fa_digits(MAX_ROWS)} ردیف داده دارد.")

            row = ImportRow(number=number, values=values)
            row.national_code = normalize_national_code(cells[columns.index(NATIONAL_CODE)])
            result.append(row)
        return result
    finally:
        workbook.close()


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------


def get_current_year():
    years = list(AcademicYear.objects.filter(is_current=True)[:2])
    if not years:
        raise FileRejected("هیچ سال تحصیلی «جاری» تعریف نشده است؛ ابتدا سال جاری را در بخش سال‌های تحصیلی مشخص کنید.")
    if len(years) > 1:
        raise FileRejected("بیش از یک سال تحصیلی «جاری» تعریف شده است؛ فقط یک سال باید جاری باشد.")
    return years[0]


def _section_key(value):
    return normalize_text(value)


def validate(rows, role, acting_user, year):
    columns = ROLE_COLUMNS[role]

    branches = {branch.code.lower(): branch for branch in Branch.objects.all()}
    allowed_branch_ids = (
        None if acting_user.is_superuser
        else set(get_accessible_branches(acting_user).values_list("id", flat=True))
    )
    grades = {grade.level: grade for grade in Grade.objects.filter(is_active=True)}
    classes = {}
    if role == STUDENT:
        for school_class in SchoolClass.objects.filter(year=year, is_active=True):
            key = (school_class.branch_id, school_class.grade_id, _section_key(school_class.section))
            classes[key] = school_class

    # Duplicates inside the file.
    rows_by_code = defaultdict(list)
    for row in rows:
        if row.national_code:
            rows_by_code[row.national_code].append(row.number)

    # People who already have an account (by national code or by username).
    codes = [code for code in rows_by_code if is_valid_national_code(code)]
    existing = {}
    for user in User.objects.filter(Q(national_code__in=codes) | Q(username__in=codes)):
        existing[user.national_code or user.username] = user
        existing.setdefault(user.username, user)
    enrolled_user_ids = set(
        StudentEnrollment.objects.filter(
            academic_year=year, student__user__in=existing.values()
        ).values_list("student__user_id", flat=True)
    )

    for row in rows:
        values = row.values
        code = row.national_code

        if code and not is_valid_national_code(code):
            row.error(f"کد ملی «{fa_digits(values[NATIONAL_CODE])}» نامعتبر است.")
        elif code and len(rows_by_code[code]) > 1:
            others = "، ".join(fa_digits(n) for n in rows_by_code[code] if n != row.number)
            row.error(f"این کد ملی در فایل تکراری است (ردیف {others}).")
        elif code and code in existing:
            # Nothing will be created for this person, so the rest of the
            # row does not matter; the existing account is left untouched.
            user = existing[code]
            if user.pk in enrolled_user_ids:
                row.skip("این دانش‌آموز در سال تحصیلی جاری ثبت‌نام دارد.")
            else:
                row.skip(f"کاربری با این کد ملی از قبل وجود دارد ({user.get_full_name() or user.username}).")
            continue

        for column in columns:
            if column not in OPTIONAL_COLUMNS and not values[column]:
                row.error(f"«{column}» خالی است.")

        _resolve_branch(row, branches, allowed_branch_ids)

        if GRADE in columns and values[GRADE]:
            level = parse_grade_level(values[GRADE])
            row.grade = grades.get(level)
            if row.grade is None:
                row.error(f"پایه «{fa_digits(values[GRADE])}» وجود ندارد یا غیرفعال است (شماره‌ی پایه را وارد کنید).")

        if role == STUDENT and row.branch and row.grade and values[SECTION]:
            row.school_class = classes.get((row.branch.id, row.grade.id, _section_key(values[SECTION])))
            if row.school_class is None:
                row.error(
                    f"کلاس «{fa_digits(values[SECTION])}» برای «{row.grade.name}» در شعبه‌ی "
                    f"«{row.branch.name}» در سال {fa_digits(year.title)} وجود ندارد یا غیرفعال است. "
                    "کلاس به‌صورت خودکار ساخته نمی‌شود."
                )

        if role == TEACHER:
            gender = parse_gender(values[GENDER])
            if gender is None:
                row.error(f"جنسیت «{values[GENDER]}» نامعتبر است (مرد/زن، پسر/دختر یا male/female).")
            else:
                row.gender = gender

    return rows


def _resolve_branch(row, branches, allowed_branch_ids):
    raw = row.values[BRANCH]
    if not raw:
        return
    branch = branches.get(raw.lower())
    if branch is None:
        row.error(f"شعبه‌ای با کد «{raw}» وجود ندارد.")
    elif not branch.is_active:
        row.error(f"شعبه‌ی «{branch.name}» غیرفعال است.")
    elif allowed_branch_ids is not None and branch.id not in allowed_branch_ids:
        row.error(f"شما به شعبه‌ی «{branch.name}» دسترسی ندارید.")
    else:
        row.branch = branch


# ---------------------------------------------------------------------------
# Creation
# ---------------------------------------------------------------------------


def _create_user(row, role):
    code = row.national_code
    user = User(
        username=code,
        national_code=code,
        first_name=row.values[FIRST_NAME],
        last_name=row.values[LAST_NAME],
        role=role,
        is_active=True,
        is_staff=False,  # students, teachers and supervisors never use the admin
        must_change_password=True,
    )
    user.set_password(code)
    user.save()
    return user


def _create_role_records(user, row, role, year, today):
    # get_or_create: reuse a profile if a signal ever starts creating them.
    if role == STUDENT:
        profile, _ = StudentProfile.objects.get_or_create(user=user, defaults={"student_code": None})
        StudentEnrollment(
            student=profile,
            school_class=row.school_class,  # save() sets academic_year from it
            enrollment_date=today,
            status=StudentEnrollment.EnrollmentStatus.ACTIVE,
        ).save()

    elif role == TEACHER:
        staff, _ = Staff.objects.get_or_create(user=user, defaults={"gender": row.gender})
        teacher, _ = TeacherProfile.objects.get_or_create(staff=staff)
        TeacherAssignment.objects.create(
            teacher=teacher,
            branch=row.branch,
            academic_year=year,
            hire_date=year.start_date,
            status=TeacherAssignment.AssignmentStatus.ACTIVE,
        )
        BranchAccess.objects.create(staff=staff, branch=row.branch, is_default=True)

    elif role == SUPERVISOR:
        SupervisorProfile.objects.get_or_create(
            user=user, defaults={"branch": row.branch, "grade": row.grade}
        )


def _log(acting_user, result, filename):
    LogEntry.objects.create(
        user_id=acting_user.pk,
        content_type_id=ContentType.objects.get_for_model(User).pk,
        object_id=None,
        object_repr=f"ورود گروهی {ROLE_LABELS[result.role]}: {fa_digits(len(result.created))} کاربر"[:200],
        action_flag=ADDITION,
        change_message=(
            f"ورود گروهی از فایل «{filename}»: {fa_digits(len(result.created))} ایجاد، "
            f"{fa_digits(len(result.skipped))} ردشده. کدهای ملی: "
            + "، ".join(user.national_code for user in result.created)
        ),
    )


def run_import(uploaded_file, role, acting_user):
    """Raises FileRejected for file-level problems; otherwise returns an ImportResult."""

    year = get_current_year()
    rows = validate(read_rows(uploaded_file, role), role, acting_user, year)
    result = ImportResult(role=role, rows=rows)

    if result.errors:
        return result  # all or nothing

    today = jdatetime.date.fromgregorian(date=timezone.localdate())
    to_create = [row for row in rows if row.status == "ok"]
    try:
        with transaction.atomic():
            for row in to_create:
                user = _create_user(row, role)
                _create_role_records(user, row, role, year, today)
                result.created.append(user)
            if result.created:
                _log(acting_user, result, getattr(uploaded_file, "name", ""))
    except IntegrityError as exc:
        # Someone created one of these accounts between validation and now.
        raise FileRejected(
            "در حین ثبت، یکی از کاربران هم‌زمان ایجاد شد و هیچ کاربری ثبت نشد. دوباره تلاش کنید."
        ) from exc

    result.committed = True
    return result
