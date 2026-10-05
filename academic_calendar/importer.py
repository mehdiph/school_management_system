"""
Bulk import of calendar events from an Excel file (admin «ورود از اکسل»).

Flow: ``read_rows`` (the raw cells) -> ``validate`` (every row, every
error, nothing written) -> the admin shows a preview -> on confirmation
``validate`` runs again on the same cells and, only when no row has an
error, ``services.import_events`` creates every event and syncs the
sessions in one transaction (all or nothing).

Columns, in this order (``COLUMNS``; the header row must match):

=================  ============================================================
عنوان               required, at most 200 characters
نوع                 «تعطیل رسمی» or «تعطیلی غیرمنتظره» (``TYPE_VALUES``)
تاریخ شروع          Jalali ``YYYY/MM/DD`` (Persian or Latin digits)
تاریخ پایان         optional, same format; defaults to the start date
کد شعبه‌ها           ``Branch.code`` values, comma separated; empty = all
پایه‌ها              grade levels (``Grade.level``, e.g. 3) or names; empty = all
زنگ‌ها               bell orders (``Bell.order``, e.g. 3) or titles; empty = all
توضیحات             optional
=================  ============================================================

Events are imported into the current academic year; both dates must be
inside it. A row that repeats an active event of that year (same title,
start and end) or another row of the file is an error.
"""

import re
from dataclasses import dataclass, field
from datetime import date, datetime

import jdatetime
from openpyxl import load_workbook

from core.services import access
from scheduling.models.bell import Bell
from school.models import AcademicYear, Branch, Grade

from .models import CalendarEvent
from .services import parse_jalali_date

TITLE = "عنوان"
TYPE = "نوع"
START = "تاریخ شروع"
END = "تاریخ پایان"
BRANCHES = "کد شعبه‌ها"
GRADES = "پایه‌ها"
BELLS = "زنگ‌ها"
DESCRIPTION = "توضیحات"

COLUMNS = [TITLE, TYPE, START, END, BRANCHES, GRADES, BELLS, DESCRIPTION]

#: Accepted «نوع» cells -> ``CalendarEvent.EventType``.
TYPE_VALUES = {
    label: value for value, label in CalendarEvent.EventType.choices
}
TYPE_ALIASES = {
    **TYPE_VALUES,
    "رسمی": CalendarEvent.EventType.OFFICIAL_HOLIDAY,
    "غیرمنتظره": CalendarEvent.EventType.UNPLANNED_CLOSURE,
    "غیر منتظره": CalendarEvent.EventType.UNPLANNED_CLOSURE,
}

MAX_ROWS = 500
MAX_TITLE = 200

_SEPARATORS = re.compile(r"[,،;؛]")
_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


class FileRejected(Exception):
    """The file as a whole cannot be read (not xlsx, wrong headers, too long...)."""


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, (datetime, date)):
        # Excel turned "1405/07/12" into a (Gregorian) date: keep its digits.
        return f"{value.year:04d}/{value.month:02d}/{value.day:02d}"
    # ZWNJ is part of Persian words; only collapse ordinary whitespace.
    return re.sub(r"[ \t\r\n]+", " ", str(value)).strip()


def _normalize(value):
    """For matching names: Arabic ي/ك -> Persian, digits -> Latin, no extra spaces."""

    return _text(value).replace("ي", "ی").replace("ك", "ک").translate(_DIGITS).strip()


def _split(value):
    return [part.strip() for part in _SEPARATORS.split(_normalize(value)) if part.strip()]


@dataclass
class ImportRow:
    number: int                 # the spreadsheet row number (header = 1)
    cells: dict                 # column -> text
    errors: list = field(default_factory=list)
    title: str = ""
    event_type: str = ""
    start_date: jdatetime.date = None
    end_date: jdatetime.date = None
    branches: list = field(default_factory=list)
    grades: list = field(default_factory=list)
    bells: list = field(default_factory=list)
    description: str = ""

    @property
    def is_valid(self):
        return not self.errors

    @property
    def type_display(self):
        """The type's label, or the cell as typed when it is not valid."""

        if self.event_type:
            return CalendarEvent.EventType(self.event_type).label
        return self.cells.get(TYPE, "")

    @property
    def start_display(self):
        return self.start_date if self.start_date else self.cells.get(START, "")

    @property
    def end_display(self):
        return self.end_date if self.end_date else self.cells.get(END, "")


@dataclass
class ImportPlan:
    academic_year: AcademicYear
    rows: list

    @property
    def errors(self):
        return [row for row in self.rows if row.errors]

    @property
    def is_valid(self):
        return bool(self.rows) and not self.errors


def read_rows(uploaded_file):
    """``[{column: text}, ...]`` from the first sheet; ``FileRejected`` when unreadable."""

    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except Exception:
        raise FileRejected("فایل قابل خواندن نیست؛ فایل اکسل (xlsx) را با قالب دانلودشده بفرستید.")

    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [_text(value) for value in next(rows, ())]
    while header and not header[-1]:
        header.pop()
    if header != COLUMNS:
        raise FileRejected(
            "ستون‌های فایل با قالب یکی نیست. ستون‌ها باید دقیقاً این‌ها باشند: "
            + "، ".join(COLUMNS)
        )

    data = []
    for values in rows:
        cells = {column: _text(value) for column, value in zip(COLUMNS, values)}
        data.append(cells)
    while data and not any(data[-1].values()):
        data.pop()
    if len(data) > MAX_ROWS:
        raise FileRejected(f"حداکثر {MAX_ROWS} ردیف در هر فایل مجاز است.")
    workbook.close()
    return data


def current_year():
    return AcademicYear.objects.filter(is_current=True).first()


def validate(raw_rows, acting_user, academic_year=None):
    """
    ``raw_rows`` (from ``read_rows``) -> ``ImportPlan``: every row parsed,
    with every problem found on it. Writes nothing.
    """

    year = academic_year or current_year()
    if year is None:
        raise FileRejected("سال تحصیلی جاری تعریف نشده است.")

    branches = {_normalize(b.code).lower(): b for b in Branch.objects.filter(is_active=True)}
    allowed_branch_ids = (
        None if acting_user.is_superuser
        else set(access.get_accessible_branch_ids(acting_user))
    )
    grades_by_level = {str(g.level): g for g in Grade.objects.filter(is_active=True)}
    grades_by_name = {_normalize(g.name): g for g in Grade.objects.filter(is_active=True)}
    bells_by_order = {str(b.order): b for b in Bell.objects.filter(is_active=True)}
    bells_by_title = {_normalize(b.title): b for b in Bell.objects.filter(is_active=True)}
    existing = set(
        CalendarEvent.objects.active()
        .filter(academic_year=year)
        .values_list("title", "start_date", "end_date")
    )

    rows, seen = [], {}
    for index, cells in enumerate(raw_rows, start=2):
        if not any(cells.values()):
            continue
        row = ImportRow(number=index, cells=cells)
        rows.append(row)

        row.title = cells.get(TITLE, "")
        if not row.title:
            row.errors.append("عنوان خالی است.")
        elif len(row.title) > MAX_TITLE:
            row.errors.append(f"عنوان نباید بیشتر از {MAX_TITLE} حرف باشد.")

        kind = _normalize(cells.get(TYPE, ""))
        row.event_type = TYPE_ALIASES.get(kind, "")
        if not row.event_type:
            row.errors.append(
                f"نوع «{cells.get(TYPE, '')}» معتبر نیست؛ یکی از این‌ها را بنویسید: "
                + "، ".join(TYPE_VALUES)
            )

        try:
            row.start_date = parse_jalali_date(cells.get(START, ""))
        except ValueError:
            row.errors.append(f"تاریخ شروع «{cells.get(START, '')}» معتبر نیست (قالب: ۱۴۰۵/۰۷/۱۲).")
        if cells.get(END):
            try:
                row.end_date = parse_jalali_date(cells[END])
            except ValueError:
                row.errors.append(f"تاریخ پایان «{cells[END]}» معتبر نیست (قالب: ۱۴۰۵/۰۷/۱۲).")
        elif row.start_date:
            row.end_date = row.start_date

        if row.start_date and row.end_date:
            if row.end_date < row.start_date:
                row.errors.append("تاریخ پایان قبل از تاریخ شروع است.")
            elif row.start_date < year.start_date or row.end_date > year.end_date:
                row.errors.append(
                    f"بازه بیرون از سال تحصیلی جاری ({year.title}: "
                    f"{year.start_date.strftime('%Y/%m/%d')} تا {year.end_date.strftime('%Y/%m/%d')}) است."
                )

        for code in _split(cells.get(BRANCHES, "")):
            branch = branches.get(code.lower())
            if branch is None:
                row.errors.append(f"شعبه‌ای با کد «{code}» وجود ندارد.")
            elif allowed_branch_ids is not None and branch.pk not in allowed_branch_ids:
                row.errors.append(f"به شعبه‌ی «{branch.name}» دسترسی ندارید.")
            else:
                row.branches.append(branch)
        if allowed_branch_ids is not None and not _split(cells.get(BRANCHES, "")):
            row.errors.append("کد شعبه را وارد کنید؛ فقط مدیر کل می‌تواند رویداد همه‌ی شعبه‌ها را ثبت کند.")

        for value in _split(cells.get(GRADES, "")):
            grade = grades_by_level.get(value) or grades_by_name.get(value)
            if grade is None:
                row.errors.append(f"پایه‌ی «{value}» پیدا نشد (شماره‌ی پایه یا نام آن را بنویسید).")
            else:
                row.grades.append(grade)

        for value in _split(cells.get(BELLS, "")):
            bell = bells_by_order.get(value) or bells_by_title.get(value)
            if bell is None:
                row.errors.append(f"زنگ «{value}» پیدا نشد (شماره‌ی ترتیب زنگ یا عنوان آن را بنویسید).")
            else:
                row.bells.append(bell)

        row.description = cells.get(DESCRIPTION, "")

        if row.title and row.start_date and row.end_date:
            key = (row.title, row.start_date, row.end_date)
            if key in existing:
                row.errors.append("رویدادی با همین عنوان و همین تاریخ‌ها از قبل ثبت شده است.")
            if key in seen:
                row.errors.append(f"تکرار ردیف {seen[key]} همین فایل است.")
            else:
                seen[key] = row.number

    if not rows:
        raise FileRejected("فایل هیچ ردیفی ندارد.")
    return ImportPlan(academic_year=year, rows=rows)
