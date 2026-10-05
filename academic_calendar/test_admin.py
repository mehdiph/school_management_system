"""
Tests for the academic calendar admin: the event form (sync after the
many-to-many scope is saved, impact message), deactivation, the quick
closure form, the conflicts page and the Excel import (template,
validation preview, all-or-nothing confirmation).
"""

from io import BytesIO
from unittest import mock

import jdatetime
from django.contrib.messages import get_messages
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from academic_calendar import importer
from academic_calendar.models import CalendarEvent
from academic_calendar.tests import SUNDAY, CalendarTestCase, Day
from core.testing import (
    grant_all_model_permissions,
    make_branch,
    make_branch_access,
    make_schedule,
    make_staff,
    make_superuser,
)
from teaching.models import SchoolSession

J = jdatetime.date


def xlsx(rows, header=importer.COLUMNS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        'events.xlsx', buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


class AdminTestCase(CalendarTestCase):

    def setUp(self):
        super().setUp()
        make_schedule(self.math, Day.SUNDAY, self.bell_1)
        make_schedule(self.math, Day.SUNDAY, self.bell_2)
        self.admin = make_superuser()
        self.client.force_login(self.admin)

    def messages(self, response):
        return [str(m) for m in get_messages(response.wsgi_request)]

    def form_data(self, **overrides):
        return {
            'academic_year': self.year.pk,
            'title': 'تاسوعا',
            'event_type': CalendarEvent.EventType.OFFICIAL_HOLIDAY,
            'start_date': '1405-07-12',
            'end_date': '1405-07-12',
            'description': '',
            'is_active': 'on',
            **overrides,
        }


class EventFormTests(AdminTestCase):

    def test_adding_an_event_cancels_sessions_and_reports_the_impact(self):
        response = self.client.post(
            reverse('admin:academic_calendar_calendarevent_add'), self.form_data(), follow=True
        )

        self.assertEqual(SchoolSession.objects.holidays().count(), 2)
        event = CalendarEvent.objects.get()
        self.assertEqual(event.created_by, self.admin)
        self.assertTrue(any('۲' in m or '2' in m for m in self.messages(response)))
        self.assertTrue(any('تداخلی پیدا نشد' in m for m in self.messages(response)))

    def test_bell_scope_saved_in_the_same_form_is_used_by_the_sync(self):
        self.client.post(
            reverse('admin:academic_calendar_calendarevent_add'),
            self.form_data(bells=[self.bell_2.pk]),
        )

        self.assertEqual(
            list(SchoolSession.objects.holidays().values_list('bell__order', flat=True)), [2]
        )

    def test_editing_the_scope_resyncs(self):
        self.client.post(reverse('admin:academic_calendar_calendarevent_add'), self.form_data())
        event = CalendarEvent.objects.get()

        self.client.post(
            reverse('admin:academic_calendar_calendarevent_change', args=[event.pk]),
            self.form_data(bells=[self.bell_1.pk]),
        )

        self.assertEqual(
            list(SchoolSession.objects.holidays().values_list('bell__order', flat=True)), [1]
        )

    def test_conflict_is_reported_after_saving(self):
        SchoolSession.objects.create(class_subject=self.math, date=SUNDAY, bell=self.bell_1)

        response = self.client.post(
            reverse('admin:academic_calendar_calendarevent_add'), self.form_data(), follow=True
        )

        self.assertTrue(any('تداخل' in m and 'conflicts' in m for m in self.messages(response)))

    def test_events_cannot_be_deleted_only_deactivated(self):
        self.client.post(reverse('admin:academic_calendar_calendarevent_add'), self.form_data())
        event = CalendarEvent.objects.get()

        response = self.client.get(reverse('admin:academic_calendar_calendarevent_delete', args=[event.pk]))
        self.assertEqual(response.status_code, 403)

        self.client.post(reverse('admin:academic_calendar_calendarevent_changelist'), {
            'action': 'deactivate_selected', '_selected_action': [event.pk],
        })
        event.refresh_from_db()
        self.assertFalse(event.is_active)
        self.assertFalse(SchoolSession.objects.exists())

    def test_branch_admin_must_pick_own_branches(self):
        staff = make_staff(is_staff=True)
        grant_all_model_permissions(staff.user)
        make_branch_access(staff, self.branch)
        self.client.force_login(staff.user)
        url = reverse('admin:academic_calendar_calendarevent_add')

        response = self.client.post(url, self.form_data())
        self.assertEqual(response.status_code, 200)  # all branches: refused
        self.assertFalse(CalendarEvent.objects.exists())

        response = self.client.post(url, self.form_data(branches=[make_branch().pk]))
        self.assertFalse(CalendarEvent.objects.exists())  # not theirs

        self.client.post(url, self.form_data(branches=[self.branch.pk]))
        self.assertEqual(CalendarEvent.objects.count(), 1)


class QuickClosureTests(AdminTestCase):

    def test_quick_closure_creates_the_event_and_cancels(self):
        response = self.client.post(reverse('admin:academic_calendar_calendarevent_quick_closure'), {
            'title': 'برف',
            'date': '1405-07-12',
            'bells': [self.bell_2.pk],
        }, follow=True)

        event = CalendarEvent.objects.get()
        self.assertEqual(event.event_type, CalendarEvent.EventType.UNPLANNED_CLOSURE)
        self.assertEqual((event.start_date, event.end_date), (SUNDAY, SUNDAY))
        self.assertEqual(SchoolSession.objects.holidays().count(), 1)
        self.assertTrue(self.messages(response))

    def test_date_outside_the_year_is_refused(self):
        response = self.client.post(reverse('admin:academic_calendar_calendarevent_quick_closure'), {
            'title': 'برف', 'date': '1404-07-12',
        })

        self.assertEqual(response.status_code, 200)
        self.assertFalse(CalendarEvent.objects.exists())


class ConflictsPageTests(AdminTestCase):

    def test_lists_held_sessions_on_closed_slots(self):
        SchoolSession.objects.create(class_subject=self.math, date=SUNDAY, bell=self.bell_1)
        self.client.post(reverse('admin:academic_calendar_calendarevent_add'), self.form_data())

        response = self.client.get(reverse('admin:academic_calendar_calendarevent_conflicts'))

        self.assertEqual(len(response.context['conflicts']), 1)
        self.assertContains(response, 'تاسوعا')
        self.assertContains(response, self.math.subject.name)


class ImportTests(AdminTestCase):

    url_name = 'admin:academic_calendar_calendarevent_import'

    def upload(self, rows, **kwargs):
        return self.client.post(reverse(self.url_name), {'file': xlsx(rows, **kwargs)})

    def confirm(self):
        return self.client.post(reverse(self.url_name), {'confirm': '1'})

    def test_template_has_the_columns_an_example_and_a_guide(self):
        response = self.client.get(reverse('admin:academic_calendar_calendarevent_import_template'))

        workbook = load_workbook(BytesIO(response.content))
        sheet = workbook.worksheets[0]
        self.assertEqual([c.value for c in sheet[1]], importer.COLUMNS)
        self.assertEqual(sheet['B2'].value, 'تعطیلی غیرمنتظره')
        self.assertIn(self.branch.code, [c.value for c in workbook['راهنما']['D']])

    def test_valid_file_previews_then_imports_and_syncs(self):
        response = self.upload([
            ['تاسوعا', 'تعطیل رسمی', '۱۴۰۵/۰۷/۱۲', '', '', '', '', ''],
            ['آلودگی', 'تعطیلی غیرمنتظره', '1405/07/19', '1405/07/20', self.branch.code,
             str(self.grade.level), '1', 'فقط زنگ اول'],
        ])

        plan = response.context['plan']
        self.assertTrue(plan.is_valid)
        self.assertFalse(CalendarEvent.objects.exists())  # preview only

        response = self.confirm()

        self.assertEqual(CalendarEvent.objects.count(), 2)
        polluted = CalendarEvent.objects.get(title='آلودگی')
        self.assertEqual(list(polluted.branches.all()), [self.branch])
        self.assertEqual(list(polluted.bells.all()), [self.bell_1])
        self.assertEqual(response.context['result'].created, 3)   # 07/12 x2, 07/19 bell 1
        self.assertEqual(SchoolSession.objects.holidays().count(), 3)

    def test_every_row_error_is_reported_and_nothing_is_imported(self):
        CalendarEvent.objects.create(
            academic_year=self.year, title='تکراری', start_date=SUNDAY, end_date=SUNDAY
        )
        response = self.upload([
            ['', 'تعطیل رسمی', '1405/07/12', '', '', '', '', ''],
            ['بد', 'نامعلوم', '1405/13/40', '', 'no-such-branch', '99', '42', ''],
            ['وارونه', 'تعطیل رسمی', '1405/07/20', '1405/07/12', '', '', '', ''],
            ['بیرون', 'تعطیل رسمی', '1404/07/12', '', '', '', '', ''],
            ['تکراری', 'تعطیل رسمی', '1405/07/12', '', '', '', '', ''],
            ['دوبار', 'تعطیل رسمی', '1405/07/13', '', '', '', '', ''],
            ['دوبار', 'تعطیل رسمی', '1405/07/13', '', '', '', '', ''],
        ])

        rows = response.context['plan'].rows
        errors = [' '.join(row.errors) for row in rows]
        self.assertIn('عنوان خالی است', errors[0])
        for text in ('نوع', 'تاریخ شروع', 'no-such-branch', '«99»', '«42»'):
            self.assertIn(text, errors[1])
        self.assertIn('قبل از تاریخ شروع', errors[2])
        self.assertIn('بیرون از سال تحصیلی', errors[3])
        self.assertIn('از قبل ثبت شده', errors[4])
        self.assertEqual(errors[5], '')
        self.assertIn('تکرار ردیف', errors[6])
        self.assertNotContains(response, 'تأیید و ثبت همه')

        # a confirm sneaked in anyway is validated again
        self.confirm()
        self.assertEqual(CalendarEvent.objects.count(), 1)

    def test_import_is_atomic_when_the_sync_fails(self):
        self.upload([['تاسوعا', 'تعطیل رسمی', '1405/07/12', '', '', '', '', '']])

        with mock.patch(
            'academic_calendar.services.sync_cancelled_sessions', side_effect=RuntimeError('boom')
        ):
            with self.assertRaises(RuntimeError), self.assertLogs('django.request', 'ERROR'):
                self.confirm()

        self.assertFalse(CalendarEvent.objects.exists())
        self.assertFalse(SchoolSession.objects.exists())

    def test_wrong_headers_are_rejected(self):
        response = self.upload([['x']], header=['a', 'b'])

        self.assertIn('ستون‌های فایل', response.context['file_error'])
