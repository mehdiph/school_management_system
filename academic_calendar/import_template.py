"""
The downloadable .xlsx template of the calendar import (see
``academic_calendar.importer`` for the column rules).

Sheet «رویدادها»: the header, one example row and Text-formatted date /
code columns (so Excel neither turns 1405/07/12 into a date nor drops
leading zeros); «نوع» gets a dropdown. Sheet «راهنما»: the accepted
types, branch codes, grade levels and bells, read from the database.
"""

from io import BytesIO

import jdatetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from core.services.access import get_accessible_branches
from scheduling.models.bell import Bell
from school.models import Branch, Grade

from .importer import (
    BELLS, BRANCHES, COLUMNS, DESCRIPTION, END, GRADES, MAX_ROWS, START, TITLE, TYPE, TYPE_VALUES,
    current_year,
)

TEXT_COLUMNS = {START, END, BRANCHES, GRADES, BELLS}
WIDTHS = {TITLE: 28, TYPE: 18, START: 14, END: 14, BRANCHES: 18, GRADES: 12, BELLS: 12, DESCRIPTION: 40}
HEADER_FILL = PatternFill("solid", fgColor="FAEAD2")


def _example_row():
    year = current_year()
    start = year.start_date if year else jdatetime.date.today()
    day = start + jdatetime.timedelta(days=10)
    return {
        TITLE: "تعطیلی به‌دلیل آلودگی هوا",
        TYPE: list(TYPE_VALUES)[1],
        START: day.strftime("%Y/%m/%d"),
        END: "",
        BRANCHES: "",
        GRADES: "",
        BELLS: "",
        DESCRIPTION: "نمونه؛ این ردیف را پیش از ارسال پاک کنید.",
    }


def _header(sheet, titles):
    for index, title in enumerate(titles, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")


def build_template(acting_user):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "رویدادها"
    sheet.sheet_view.rightToLeft = True
    sheet.freeze_panes = "A2"

    _header(sheet, COLUMNS)
    for index, title in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = WIDTHS.get(title, 16)
        if title in TEXT_COLUMNS:
            for row in range(2, MAX_ROWS + 2):
                sheet.cell(row=row, column=index).number_format = "@"

    example = _example_row()
    for index, title in enumerate(COLUMNS, start=1):
        sheet.cell(row=2, column=index, value=example[title]).font = Font(italic=True, color="808080")

    letter = get_column_letter(COLUMNS.index(TYPE) + 1)
    validation = DataValidation(
        type="list",
        formula1='"' + ",".join(TYPE_VALUES) + '"',
        allow_blank=False,
        showErrorMessage=True,
        errorTitle="مقدار نامعتبر",
        error="«نوع» را از فهرست انتخاب کنید.",
    )
    validation.add(f"{letter}2:{letter}{MAX_ROWS + 1}")
    sheet.add_data_validation(validation)

    guide = workbook.create_sheet("راهنما")
    guide.sheet_view.rightToLeft = True
    branches = (
        Branch.objects.filter(is_active=True)
        if acting_user.is_superuser
        else get_accessible_branches(acting_user).filter(is_active=True)
    ).order_by("order", "name")
    sections = [
        ("نوع", [(label, "") for label in TYPE_VALUES]),
        ("کد شعبه", [(b.code, b.name) for b in branches]),
        ("پایه (شماره)", [(g.level, g.name) for g in Grade.objects.filter(is_active=True).order_by("level")]),
        ("زنگ (ترتیب)", [
            (b.order, f"{b.title} ({b.start_time:%H:%M}–{b.end_time:%H:%M})")
            for b in Bell.objects.filter(is_active=True).order_by("order")
        ]),
    ]
    column = 1
    for title, values in sections:
        guide.cell(row=1, column=column, value=title)
        guide.cell(row=1, column=column).font = Font(bold=True)
        guide.cell(row=1, column=column).fill = HEADER_FILL
        guide.cell(row=1, column=column + 1, value="شرح").font = Font(bold=True)
        guide.cell(row=1, column=column + 1).fill = HEADER_FILL
        for row, (value, label) in enumerate(values, start=2):
            guide.cell(row=row, column=column, value=value)
            guide.cell(row=row, column=column + 1, value=label)
        guide.column_dimensions[get_column_letter(column)].width = 18
        guide.column_dimensions[get_column_letter(column + 1)].width = 26
        column += 3

    notes = [
        "تاریخ‌ها شمسی و به شکل ۱۴۰۵/۰۷/۱۲ هستند؛ تاریخ پایان اختیاری است و خالی یعنی همان روز شروع.",
        "شعبه، پایه و زنگ را با ویرگول جدا کنید؛ خالی یعنی «همه».",
        "پنج‌شنبه و جمعه همیشه تعطیل‌اند و جزو تقویم نیستند؛ اگر بازه آن‌ها را بپوشاند نادیده گرفته می‌شوند.",
        "اگر حتی یک ردیف خطا داشته باشد، هیچ رویدادی ثبت نمی‌شود.",
    ]
    start = max(len(values) for _, values in sections) + 3
    for offset, note in enumerate(notes):
        guide.cell(row=start + offset, column=1, value=note)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
